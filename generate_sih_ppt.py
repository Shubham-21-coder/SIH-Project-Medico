import os
import sys
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_high_quality_sih_deck(output_path="Codestreaker_AyushSetu_SIH2026.pptx"):
    # Source image paths
    brain_dir = r"C:\Users\asus\.gemini\antigravity\brain\b7835552-5eda-452b-bc4d-59f7cacaf3f2"
    img1_src = os.path.join(brain_dir, "patient_kiosk_intake_1788059711956.jpg")
    img2_src = os.path.join(brain_dir, "smart_clinical_engine_1788059732903.jpg")
    img3_src = os.path.join(brain_dir, "doctor_dashboard_briefing_1788059751315.jpg")

    assets_dir = os.path.join(os.getcwd(), "assets")
    os.makedirs(assets_dir, exist_ok=True)
    
    img1_path = os.path.join(assets_dir, "patient_kiosk.jpg")
    img2_path = os.path.join(assets_dir, "clinical_engine.jpg")
    img3_path = os.path.join(assets_dir, "doctor_dashboard.jpg")

    if os.path.exists(img1_src): shutil.copy2(img1_src, img1_path)
    if os.path.exists(img2_src): shutil.copy2(img2_src, img2_path)
    if os.path.exists(img3_src): shutil.copy2(img3_src, img3_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    NAVY_DARK = RGBColor(15, 33, 64)       # #0F2140
    NAVY_PRIMARY = RGBColor(29, 112, 184)   # #1D70B8
    BLUE_LIGHT = RGBColor(238, 246, 255)    # #EEF6FF
    TEAL_ACCENT = RGBColor(0, 168, 132)     # #00A884
    TEAL_LIGHT = RGBColor(230, 250, 245)    # #E6FAF5
    AMBER_ACCENT = RGBColor(217, 119, 6)    # #D97706
    RED_ACCENT = RGBColor(225, 29, 72)      # #E11D48
    BG_LIGHT = RGBColor(248, 250, 252)      # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF
    CARD_BORDER = RGBColor(218, 226, 237)   # #DAE2ED
    TEXT_DARK = RGBColor(15, 23, 42)        # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
    WHITE = RGBColor(255, 255, 255)

    def add_top_team_badge(slide, team_name="Codestreaker"):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.2), Inches(1.9), Inches(0.7))
        badge.fill.solid()
        badge.fill.fore_color.rgb = WHITE
        badge.line.color.rgb = NAVY_PRIMARY
        badge.line.width = Pt(1.5)
        tf = badge.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = "Team:"
        p1.font.size = Pt(8.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.font.name = "Arial"
        p1.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = team_name
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = NAVY_DARK
        p2.font.name = "Arial"
        p2.alignment = PP_ALIGN.CENTER

    def add_top_sih_header(slide, slide_title=""):
        sih_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.2), Inches(2.4), Inches(0.7))
        tf_sih = sih_box.text_frame
        p_sih = tf_sih.paragraphs[0]
        p_sih.text = "SMART INDIA\nHACKATHON 2026"
        p_sih.font.size = Pt(9.5)
        p_sih.font.bold = True
        p_sih.font.color.rgb = NAVY_PRIMARY
        p_sih.font.name = "Arial"
        p_sih.alignment = PP_ALIGN.RIGHT

        if slide_title:
            title_box = slide.shapes.add_textbox(Inches(2.4), Inches(0.18), Inches(8.0), Inches(0.75))
            tf_title = title_box.text_frame
            p_t = tf_title.paragraphs[0]
            p_t.text = slide_title
            p_t.font.size = Pt(19)
            p_t.font.bold = True
            p_t.font.color.rgb = NAVY_DARK
            p_t.font.name = "Arial"
            p_t.alignment = PP_ALIGN.CENTER

    def add_bottom_footer(slide, slide_num=2):
        footer_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
        footer_bg.fill.solid()
        footer_bg.fill.fore_color.rgb = NAVY_PRIMARY
        footer_bg.line.fill.background()

        footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(7.08), Inches(12.333), Inches(0.4))
        tf_f = footer_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = f"@SIH Idea submission- Template                                                                                                                {slide_num}"
        p_f.font.size = Pt(10)
        p_f.font.color.rgb = WHITE
        p_f.font.name = "Arial"

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_LIGHT
    bg1.line.fill.background()

    top_bar = s1.shapes.add_textbox(Inches(1.0), Inches(0.4), Inches(11.333), Inches(0.8))
    tf_top = top_bar.text_frame
    p_top = tf_top.paragraphs[0]
    p_top.text = "SMART INDIA HACKATHON 2026"
    p_top.font.size = Pt(26)
    p_top.font.bold = True
    p_top.font.color.rgb = NAVY_PRIMARY
    p_top.font.name = "Arial"
    p_top.alignment = PP_ALIGN.CENTER

    p_sub = tf_top.add_paragraph()
    p_sub.text = "TITLE PAGE"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = NAVY_DARK
    p_sub.font.name = "Arial"
    p_sub.alignment = PP_ALIGN.CENTER

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.45), Inches(11.333), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    c1.line.width = Pt(1.5)

    info_box = s1.shapes.add_textbox(Inches(1.4), Inches(1.65), Inches(10.5), Inches(5.0))
    tf_info = info_box.text_frame
    tf_info.word_wrap = True

    fields = [
        ("• Problem Statement ID –", "[Enter Problem Statement ID / Student Innovation]"),
        ("• Problem Statement Title –", "AI-Powered Pre-Consultation Clinical Intake, OPD Triage & ABDM Health Kiosk"),
        ("• Theme –", "MedTech / Healthcare / Smart Health & AYUSH Integration"),
        ("• PS Category –", "Software"),
        ("• Team ID –", "[Enter Team ID from SIH Portal]"),
        ("• Team Name (Registered on portal) –", "Codestreaker"),
        ("• Project Name –", "Ayush Setu (आयुष सेतु) — Pre-Consultation Clinical Intelligence System"),
    ]

    for i, (label, val) in enumerate(fields):
        p = tf_info.paragraphs[0] if i == 0 else tf_info.add_paragraph()
        p.space_after = Pt(10)
        p.space_before = Pt(4)
        run_l = p.add_run()
        run_l.text = label + " "
        run_l.font.bold = True
        run_l.font.size = Pt(14)
        run_l.font.color.rgb = NAVY_PRIMARY
        run_l.font.name = "Arial"

        run_v = p.add_run()
        run_v.text = val
        run_v.font.bold = (i in [5, 6])
        run_v.font.size = Pt(14)
        run_v.font.color.rgb = TEAL_ACCENT if i == 6 else (NAVY_DARK if i == 5 else TEXT_DARK)
        run_v.font.name = "Arial"

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION (HIGH-QUALITY VISUAL STORYBOARD)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_top_team_badge(s2, "Codestreaker")
    add_top_sih_header(s2, "AYUSH SETU — SMART CLINICAL INTAKE KIOSK")
    add_bottom_footer(s2, 2)

    # Sub-heading banner
    sub2 = s2.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.333), Inches(0.4))
    p2_sub = sub2.text_frame.paragraphs[0]
    p2_sub.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype) — End-to-End Visual Journey"
    p2_sub.font.size = Pt(14)
    p2_sub.font.bold = True
    p2_sub.font.color.rgb = NAVY_PRIMARY
    p2_sub.font.name = "Arial"

    # 3 STAGES LAYOUT
    col_w = Inches(3.7)
    col_h = Inches(4.35)
    col_y = Inches(1.35)

    # --- STAGE 01: PATIENT INTAKE ---
    c1_x = Inches(0.5)
    card1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_x, col_y, col_w, col_h)
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = NAVY_PRIMARY
    card1.line.width = Pt(1.5)

    # Stage 01 Header Badge
    h1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_x + Inches(0.2), col_y + Inches(0.12), col_w - Inches(0.4), Inches(0.38))
    h1.fill.solid()
    h1.fill.fore_color.rgb = BLUE_LIGHT
    h1.line.color.rgb = NAVY_PRIMARY
    h1.line.width = Pt(1)
    tf_h1 = h1.text_frame
    p_h1 = tf_h1.paragraphs[0]
    p_h1.text = "STAGE 01: PATIENT INTAKE"
    p_h1.font.size = Pt(11)
    p_h1.font.bold = True
    p_h1.font.color.rgb = NAVY_PRIMARY
    p_h1.alignment = PP_ALIGN.CENTER

    # Image 1
    if os.path.exists(img1_path):
        s2.shapes.add_picture(img1_path, c1_x + Inches(0.85), col_y + Inches(0.58), width=Inches(2.0), height=Inches(2.0))

    # Stage 01 Text Content
    t1_box = s2.shapes.add_textbox(c1_x + Inches(0.15), col_y + Inches(2.65), col_w - Inches(0.3), Inches(1.6))
    tf_t1 = t1_box.text_frame
    tf_t1.word_wrap = True

    p = tf_t1.paragraphs[0]
    p.text = "Capture Before Consultation"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = NAVY_DARK
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)

    p = tf_t1.add_paragraph()
    p.text = "Patient shares symptoms, medical history & basic vitals at the smart kiosk before meeting the doctor."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)

    # Tags Row
    p_tags = tf_t1.add_paragraph()
    p_tags.alignment = PP_ALIGN.CENTER
    r1 = p_tags.add_run()
    r1.text = "🎤 Voice + Text   "
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = NAVY_PRIMARY

    r2 = p_tags.add_run()
    r2.text = "🌐 8 Indian Languages   "
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = TEAL_ACCENT

    r3 = p_tags.add_run()
    r3.text = "📋 Guided Intake"
    r3.font.bold = True
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = NAVY_DARK

    # --- CONNECTING ARROW 1 ---
    arrow1 = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.25), col_y + Inches(1.8), Inches(0.55), Inches(0.35))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = NAVY_PRIMARY
    arrow1.line.fill.background()

    # --- STAGE 02: SMART CLINICAL ENGINE (HERO) ---
    c2_x = Inches(4.85)
    card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_x, col_y, col_w, col_h)
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = TEAL_ACCENT
    card2.line.width = Pt(2.0)

    # Stage 02 Header Badge
    h2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_x + Inches(0.2), col_y + Inches(0.12), col_w - Inches(0.4), Inches(0.38))
    h2.fill.solid()
    h2.fill.fore_color.rgb = TEAL_LIGHT
    h2.line.color.rgb = TEAL_ACCENT
    h2.line.width = Pt(1)
    tf_h2 = h2.text_frame
    p_h2 = tf_h2.paragraphs[0]
    p_h2.text = "STAGE 02: SMART CLINICAL ENGINE"
    p_h2.font.size = Pt(11)
    p_h2.font.bold = True
    p_h2.font.color.rgb = TEAL_ACCENT
    p_h2.alignment = PP_ALIGN.CENTER

    # Image 2
    if os.path.exists(img2_path):
        s2.shapes.add_picture(img2_path, c2_x + Inches(0.85), col_y + Inches(0.58), width=Inches(2.0), height=Inches(2.0))

    # Stage 02 Text Content
    t2_box = s2.shapes.add_textbox(c2_x + Inches(0.15), col_y + Inches(2.65), col_w - Inches(0.3), Inches(1.6))
    tf_t2 = t2_box.text_frame
    tf_t2.word_wrap = True

    p = tf_t2.paragraphs[0]
    p.text = "Structure & Synthesize"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = NAVY_DARK
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)

    p = tf_t2.add_paragraph()
    p.text = "Transforms raw multilingual answers into a concise, standardized doctor brief & HL7 FHIR R4 record."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)

    # Tags Row
    p_tags2 = tf_t2.add_paragraph()
    p_tags2.alignment = PP_ALIGN.CENTER
    badges = [
        ("[Chief Complaint]", NAVY_PRIMARY),
        ("[HPI]", NAVY_DARK),
        ("[Vitals]", TEAL_ACCENT),
        ("[🚨 Red Flags]", RED_ACCENT),
    ]
    for b_txt, b_col in badges:
        r = p_tags2.add_run()
        r.text = b_txt + " "
        r.font.bold = True
        r.font.size = Pt(8.0)
        r.font.color.rgb = b_col

    # --- CONNECTING ARROW 2 ---
    arrow2 = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.6), col_y + Inches(1.8), Inches(0.55), Inches(0.35))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = TEAL_ACCENT
    arrow2.line.fill.background()

    # --- STAGE 03: DOCTOR BRIEFING ---
    c3_x = Inches(9.2)
    card3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_x, col_y, col_w, col_h)
    card3.fill.solid()
    card3.fill.fore_color.rgb = CARD_BG
    card3.line.color.rgb = NAVY_PRIMARY
    card3.line.width = Pt(1.5)

    # Stage 03 Header Badge
    h3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_x + Inches(0.2), col_y + Inches(0.12), col_w - Inches(0.4), Inches(0.38))
    h3.fill.solid()
    h3.fill.fore_color.rgb = BLUE_LIGHT
    h3.line.color.rgb = NAVY_PRIMARY
    h3.line.width = Pt(1)
    tf_h3 = h3.text_frame
    p_h3 = tf_h3.paragraphs[0]
    p_h3.text = "STAGE 03: DOCTOR BRIEFING"
    p_h3.font.size = Pt(11)
    p_h3.font.bold = True
    p_h3.font.color.rgb = NAVY_PRIMARY
    p_h3.alignment = PP_ALIGN.CENTER

    # Image 3
    if os.path.exists(img3_path):
        s2.shapes.add_picture(img3_path, c3_x + Inches(0.85), col_y + Inches(0.58), width=Inches(2.0), height=Inches(2.0))

    # Stage 03 Text Content
    t3_box = s2.shapes.add_textbox(c3_x + Inches(0.15), col_y + Inches(2.65), col_w - Inches(0.3), Inches(1.6))
    tf_t3 = t3_box.text_frame
    tf_t3.word_wrap = True

    p = tf_t3.paragraphs[0]
    p.text = "Doctor Sees Context First"
    p.font.bold = True
    p.font.size = Pt(11.5)
    p.font.color.rgb = NAVY_DARK
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)

    p = tf_t3.add_paragraph()
    p.text = "Physician reviews structured patient history beforehand, eliminating repetitive questions."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)

    # Tags Row
    p_tags3 = tf_t3.add_paragraph()
    p_tags3.alignment = PP_ALIGN.CENTER
    r1 = p_tags3.add_run()
    r1.text = "⚡ Faster Review   "
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = NAVY_PRIMARY

    r2 = p_tags3.add_run()
    r2.text = "🔍 Upfront Context   "
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = TEAL_ACCENT

    r3 = p_tags3.add_run()
    r3.text = "🚨 Red-Flag Alert"
    r3.font.bold = True
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = RED_ACCENT

    # --- BOTTOM INNOVATION STRIP: WHAT MAKES AYUSH SETU DIFFERENT ---
    strip_y = Inches(5.8)
    strip_h = Inches(1.1)
    strip_w = Inches(12.4)
    
    strip_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), strip_y, strip_w, strip_h)
    strip_card.fill.solid()
    strip_card.fill.fore_color.rgb = BG_LIGHT
    strip_card.line.color.rgb = CARD_BORDER
    strip_card.line.width = Pt(1.5)

    # Title of Strip
    strip_lbl = s2.shapes.add_textbox(Inches(0.6), strip_y + Inches(0.04), Inches(3.0), Inches(0.3))
    p_sl = strip_lbl.text_frame.paragraphs[0]
    p_sl.text = "★ WHAT MAKES AYUSH SETU DIFFERENT:"
    p_sl.font.size = Pt(9)
    p_sl.font.bold = True
    p_sl.font.color.rgb = NAVY_DARK

    # 4 Innovation Highlights
    pill_w = Inches(2.8)
    pill_h = Inches(0.65)
    pill_y = strip_y + Inches(0.35)

    pills_data = [
        (Inches(0.7), "🌐 8 INDIAN LANGUAGES", "Voice & text with zero question repetition", NAVY_PRIMARY),
        (Inches(3.8), "🩺 ALLOPATHY + AYUSH", "Dual modern clinical & Dashavidha intake", TEAL_ACCENT),
        (Inches(6.9), "🔐 ABDM / ABHA READY", "14-digit ABHA ID & DPDP Act consent", NAVY_DARK),
        (Inches(10.0), "🚨 RED-FLAG HIGHLIGHTING", "Instant prioritization of critical symptoms", RED_ACCENT),
    ]

    for px, ptitle, pdesc, pcol in pills_data:
        pill = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, pill_y, pill_w, pill_h)
        pill.fill.solid()
        pill.fill.fore_color.rgb = WHITE
        pill.line.color.rgb = pcol
        pill.line.width = Pt(1.2)
        tf_p = pill.text_frame
        tf_p.word_wrap = True
        
        p = tf_p.paragraphs[0]
        p.text = ptitle
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = pcol
        p.alignment = PP_ALIGN.CENTER
        
        p2 = tf_p.add_paragraph()
        p2.text = pdesc
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_top_team_badge(s3, "Codestreaker")
    add_top_sih_header(s3, "TECHNICAL APPROACH & METHODOLOGY")
    add_bottom_footer(s3, 3)

    box_w_l = Inches(5.8)
    c3_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.2), box_w_l, Inches(5.6))
    c3_l.fill.solid()
    c3_l.fill.fore_color.rgb = CARD_BG
    c3_l.line.color.rgb = CARD_BORDER
    tf3_l = c3_l.text_frame
    tf3_l.word_wrap = True

    p = tf3_l.paragraphs[0]
    p.text = "Technologies Used & Architecture Stack"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(6)

    techs = [
        ("Frontend (React 19, TypeScript, Vite)", "Used to build responsive, accessible patient kiosk & physician portals with glassmorphism UI."),
        ("Backend (Node.js 24, Express 4, TypeScript)", "Handles modular REST API routing, session management, and clinical orchestration."),
        ("Session Storage & Data Layer", "In-memory UUID clinical context store (`server/sessions.ts`) + JSON persistence for users and sessions."),
        ("Authentication & Verification", "Official ABDM M1 Gateway (`/v0.5/sessions`), Nodemailer (Gmail SMTP), Twilio Verify SMS with SHA-256 Hashing."),
        ("AI / Clinical Intelligence", "Google Gemini AI (`@google/generative-ai`) with deterministic sequential clinical triage rules & FHIR R4 JSON generator."),
        ("Accessibility & Audio Engine", "Web Speech API (SpeechSynthesis) for Indic voice consent + Web Audio API for interactive kiosk chimes."),
    ]

    for t_name, t_desc in techs:
        p = tf3_l.add_paragraph()
        run_t = p.add_run()
        run_t.text = "▪ " + t_name + ": "
        run_t.font.bold = True
        run_t.font.size = Pt(9.5)
        run_t.font.color.rgb = NAVY_DARK
        run_d = p.add_run()
        run_d.text = t_desc
        run_d.font.size = Pt(9.5)
        run_d.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    box_w_r = Inches(6.0)
    c3_r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.2), box_w_r, Inches(5.6))
    c3_r.fill.solid()
    c3_r.fill.fore_color.rgb = CARD_BG
    c3_r.line.color.rgb = CARD_BORDER
    tf3_r = c3_r.text_frame
    tf3_r.word_wrap = True

    p = tf3_r.paragraphs[0]
    p.text = "Methodology & End-to-End Implementation Workflow"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = TEAL_ACCENT
    p.space_after = Pt(6)

    workflow_steps = [
        ("Step 1: Patient Check-in & Consent", "Patient registers via 14-Digit ABHA ID, Phone, or Email; verifies identity with real OTP; grants DPDP Act 2023 consent with spoken voice feedback."),
        ("Step 2: Multilingual Adaptive Triage", "Sequential clinical questioning across 8 Indic languages for cardiology, respiratory, GI, fever, or AYUSH Dashavidha without question repetition."),
        ("Step 3: Document Digitization & Linking", "Patient uploads prior paper prescriptions/lab reports; system digitizes metadata and links with ABHA profile."),
        ("Step 4: Clinical Summarization & FHIR R4", "Engine constructs structured Chief Complaint, HPI chronology, red flags, and HL7 FHIR R4 Bundle ready for hospital EHR transmission."),
        ("Step 5: Doctor Pre-Consultation Review", "Treating physician reviews structured patient briefing on Doctor Portal BEFORE patient enters, enabling highly focused, efficient clinical care."),
    ]

    for s_title, s_desc in workflow_steps:
        p = tf3_r.add_paragraph()
        run_st = p.add_run()
        run_st.text = "➔ " + s_title + "\n"
        run_st.font.bold = True
        run_st.font.size = Pt(9.5)
        run_st.font.color.rgb = NAVY_PRIMARY
        run_sd = p.add_run()
        run_sd.text = "   " + s_desc
        run_sd.font.size = Pt(9)
        run_sd.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_top_team_badge(s4, "Codestreaker")
    add_top_sih_header(s4, "FEASIBILITY AND VIABILITY")
    add_bottom_footer(s4, 4)

    c4_w = Inches(3.8)
    c4_h = Inches(5.4)
    c4_y = Inches(1.35)

    c4_1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), c4_y, c4_w, c4_h)
    c4_1.fill.solid()
    c4_1.fill.fore_color.rgb = CARD_BG
    c4_1.line.color.rgb = CARD_BORDER
    tf4_1 = c4_1.text_frame
    tf4_1.word_wrap = True

    p = tf4_1.paragraphs[0]
    p.text = "Analysis of Feasibility"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY_DARK
    p.space_after = Pt(8)

    feas_points = [
        "Low-Cost Hardware Deployment: Runs seamlessly on standard touch tablets, Android kiosks, desktop computers, or patient smartphones via QR codes.",
        "Zero High-End Compute Barrier: Edge-ready web architecture with lightweight REST APIs, running smoothly without requiring expensive GPUs at clinic sites.",
        "Government Standards Compliant: Fully aligned with Ayushman Bharat Digital Mission (ABDM) M1 standards and DPDP Act 2023 legal consent frameworks.",
        "Immediate Hospital Adoption: Plug-and-play browser interface requiring zero complex server installations for government OPD clinics.",
    ]
    for pt in feas_points:
        p = tf4_1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    c4_2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.75), c4_y, c4_w, c4_h)
    c4_2.fill.solid()
    c4_2.fill.fore_color.rgb = CARD_BG
    c4_2.line.color.rgb = CARD_BORDER
    tf4_2 = c4_2.text_frame
    tf4_2.word_wrap = True

    p = tf4_2.paragraphs[0]
    p.text = "Potential Challenges & Risks"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = AMBER_ACCENT
    p.space_after = Pt(8)

    risks_points = [
        "Digital Literacy in Rural Patients: Elderly or illiterate patients may struggle with text inputs on digital touchscreens.",
        "Peak OPD High-Concurrency: Hundreds of patients arriving simultaneously at government hospital registration counters during morning rush hours.",
        "Telecom Carrier SMS Delays: Network congestion in remote rural areas delaying telecom SMS OTP delivery.",
        "Data Privacy & Compliance: Risk of sensitive health history exposure without stringent cryptographic safeguards.",
    ]
    for pt in risks_points:
        p = tf4_2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    c4_3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), c4_y, c4_w, c4_h)
    c4_3.fill.solid()
    c4_3.fill.fore_color.rgb = CARD_BG
    c4_3.line.color.rgb = CARD_BORDER
    tf4_3 = c4_3.text_frame
    tf4_3.word_wrap = True

    p = tf4_3.paragraphs[0]
    p.text = "Strategies for Overcoming Challenges"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = TEAL_ACCENT
    p.space_after = Pt(8)

    mit_points = [
        "Multilingual Voice Readouts: Web Speech API audio guidance in 8 Indic languages + Staff-Assisted check-in mode for non-tech-savvy patients.",
        "Lightweight In-Memory Caching: High-throughput Node.js microservices responding in <50ms, easily handling peak hospital loads.",
        "Multi-Channel Dual OTP: Dual telecom routing (Twilio / Fast2SMS / 2Factor) + instant Nodemailer Email OTP fallback ensures zero patient blocking.",
        "End-to-End Security: SHA-256 OTP hashing with server salt, 256-bit session tokens, and DPDP Act 2023 digital consent logs.",
    ]
    for pt in mit_points:
        p = tf4_3.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_top_team_badge(s5, "Codestreaker")
    add_top_sih_header(s5, "IMPACT AND BENEFITS")
    add_bottom_footer(s5, 5)

    c5_top = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.133), Inches(2.6))
    c5_top.fill.solid()
    c5_top.fill.fore_color.rgb = CARD_BG
    c5_top.line.color.rgb = CARD_BORDER
    tf5_t = c5_top.text_frame
    tf5_t.word_wrap = True

    p = tf5_t.paragraphs[0]
    p.text = "Potential Impact on Target Audience"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(6)

    aud_points = [
        "For Patients: Zero intimidation in clinic visits; freedom to express symptoms in native mother tongue (Hindi, Tamil, Marathi, etc.); drastic reduction in overall waiting room anxiety and consultation delays.",
        "For Doctors & Clinicians: Pre-prepared, structured clinical briefing available on screen before the patient walks into the cabin. Zero time wasted repeating basic demographic and chronological symptom questions.",
        "For Hospital Administrators: Significant optimization of outpatient queuing; structured digital records conforming to Ayushman Bharat Digital Mission (ABDM) standards; enhanced hospital operational throughput.",
    ]
    for pt in aud_points:
        p = tf5_t.add_paragraph()
        p.text = "▪ " + pt
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    c5_bot = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.1), Inches(12.133), Inches(2.7))
    c5_bot.fill.solid()
    c5_bot.fill.fore_color.rgb = CARD_BG
    c5_bot.line.color.rgb = CARD_BORDER
    tf5_b = c5_bot.text_frame
    tf5_b.word_wrap = True

    p = tf5_b.paragraphs[0]
    p.text = "Multidimensional Benefits of the Solution"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = TEAL_ACCENT
    p.space_after = Pt(6)

    ben_points = [
        ("Social Impact: ", "Democratizes high-quality clinical intake for underserved and non-English-speaking populations across urban and rural India, bridging the gap between traditional AYUSH practices and modern healthcare systems."),
        ("Economic Impact: ", "Saves millions of clinical man-hours annually in government and private hospitals; reduces administrative charting expenses; increases daily patient intake capacity per doctor by up to 40%."),
        ("Environmental Impact: ", "Enables 100% paperless OPD intake, eliminating physical paper case sheets, manual token slips, and duplicated prescription records through cloud/local FHIR R4 digital exchange."),
    ]
    for b_title, b_desc in ben_points:
        p = tf5_b.add_paragraph()
        run_bt = p.add_run()
        run_bt.text = "✔ " + b_title
        run_bt.font.bold = True
        run_bt.font.size = Pt(10)
        run_bt.font.color.rgb = NAVY_DARK
        run_bd = p.add_run()
        run_bd.text = b_desc
        run_bd.font.size = Pt(10)
        run_bd.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_top_team_badge(s6, "Codestreaker")
    add_top_sih_header(s6, "RESEARCH AND REFERENCES")
    add_bottom_footer(s6, 6)

    c6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.133), Inches(5.5))
    c6.fill.solid()
    c6.fill.fore_color.rgb = CARD_BG
    c6.line.color.rgb = CARD_BORDER
    tf6 = c6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "Details / Links of Reference and Research Work"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(10)

    refs = [
        ("National Health Authority (NHA) & ABDM Architecture Specification", "Guidelines on Milestone 1 (M1) ABHA Verification, Health Facility Registry (HFR), and PHR Handle Authentication.\nhttps://sandbox.abdm.gov.in/ | https://abdm.gov.in/"),
        ("HL7 FHIR R4 (Fast Healthcare Interoperability Resources) Standard", "FHIR Document Bundle specifications for Electronic Health Records (EHR) transmission and clinical diagnostic exchange.\nhttps://hl7.org/fhir/R4/"),
        ("Digital Personal Data Protection (DPDP) Act 2023 Compliance Framework", "Government of India statutory guidelines on informed, purpose-bound patient consent, data minimization, and audit trails for digital health platforms.\nhttps://www.meity.gov.in/data-protection-framework"),
        ("Ministry of AYUSH — Dashavidha & Ashtavidha Pariksha Clinical Guidelines", "Traditional Indian systems of medicine clinical examination guidelines for Prakriti, Agni, Koshtha, and Dhatu imbalance triage.\nhttps://ayush.gov.in/"),
        ("Indian Public Health Standards (IPHS) Guidelines for OPD Queuing & Clinical Consultation", "Research on outpatient department wait-time reduction, clinical history-taking time allocation, and physician burnout mitigation.\nhttps://nhsrcindia.org/indian-public-health-standards"),
    ]

    for idx, (r_title, r_desc) in enumerate(refs, start=1):
        p = tf6.add_paragraph()
        run_rt = p.add_run()
        run_rt.text = f"[{idx}] {r_title}\n"
        run_rt.font.bold = True
        run_rt.font.size = Pt(10.5)
        run_rt.font.color.rgb = NAVY_DARK
        run_rd = p.add_run()
        run_rd.text = f"     {r_desc}"
        run_rd.font.size = Pt(9.5)
        run_rd.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    try:
        prs.save(output_path)
        print(f"SUCCESS: High-Quality SIH Presentation generated at {output_path}")
    except Exception as e:
        fallback_path = "Codestreaker_AyushSetu_SIH2026_Final.pptx"
        prs.save(fallback_path)
        print(f"SUCCESS: Presentation saved at {fallback_path} (Original file was locked by PowerPoint)")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "Codestreaker_AyushSetu_SIH2026_Final.pptx"
    build_high_quality_sih_deck(out_file)

